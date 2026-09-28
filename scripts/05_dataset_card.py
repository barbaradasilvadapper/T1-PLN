"""Etapa 5 (item i): dataset card.

Preenche o template da disciplina (docs/templateDatasetCard.xlsx) com os valores do corpus e grava
DATASET_CARD.xlsx (mesmo layout; a coluna "Exemplo" vira "Valor") e DATASET_CARD.md na raiz de T1.
Todos os números vêm de corpus/estatisticas/estatisticas.json; a validação da classificação vem de
corpus/validacao_manual.json (se existir).
"""
import json
from copy import copy
from datetime import date

from openpyxl import load_workbook
from openpyxl.styles import Alignment

from comum import CORPUS, RAIZ

CRIADOR = "Bárbara da Silva Dapper (barbaradasilvadapper@gmail.com)"
LOCAL = "Pasta do trabalho: T1/corpus/corpus.json (URL de hospedagem - GitHub ou Hugging Face - a definir)"
VERSAO = "1.0"


def tamanho(p):
    b = p.stat().st_size
    return f"{b / 1024 / 1024:.1f} MB" if b > 1024 * 1024 else f"{b / 1024:.0f} KB"


def valores():
    E = json.load(open(CORPUS / "estatisticas" / "estatisticas.json", encoding="utf-8"))
    meta = json.load(open(CORPUS / "corpus.json", encoding="utf-8"))["metadados"]
    g = E["geral"]
    hoje = date.fromisoformat(meta["data"]).strftime("%d/%m/%Y")
    anos = list(E["por_ano"])
    dist = "; ".join(f"{s['nome']}: {s['questoes']} questões ({str(s['percentual']).replace('.', ',')}%)"
                     for s in E["por_subarea"].values())
    val = None
    vpath = CORPUS / "validacao_manual.json"
    if vpath.exists():
        val = json.load(open(vpath, encoding="utf-8"))
    qual = E.get("qualidade_lexica")
    outras = [
        "Banca: FGV (todas as provas), tipo de caderno 1; gabarito oficial (definitivo quando disponível no site, "
        "senão preliminar); questões anuladas foram removidas.",
        "Classes (subáreas) atribuídas por um classificador de palavras-chave ponderadas (scripts/subareas.py), "
        "não pelas seções da prova, porque uma mesma prova mistura vários assuntos.",
    ]
    if val:
        outras.append(f"Conferência manual de uma amostra aleatória estratificada de {val['n']} questões (feita com auxílio "
                      f"do Claude, IA): {val['dominio_correto']} de {val['n']} são de computação e "
                      f"{val['subarea_correta']} de {val['n']} estão na subárea correta "
                      f"({val['subarea_correta'] / val['n']:.0%}).")
    outras += [
        f"Pré-processamento: conversão PDF->texto com poppler respeitando as 2 colunas; separação de enunciado e "
        f"alternativas; remoção de {g['questoes_descartadas']} das {g['questoes_extraidas_dos_pdfs']} questões extraídas "
        "(outras matérias, problemas de conversão como fórmulas/tabelas quebradas ou dependência de figura, "
        "sem gabarito, anuladas e duplicatas entre provas do mesmo concurso).",
    ]
    if qual:
        outras.append(f"Qualidade dos textos = proporção de tokens alfabéticos reconhecidos em dicionário de português "
                      f"ou inglês (pyspellchecker): {qual['proporcao_reconhecida']:.1%}. É um limite inferior: os não "
                      "reconhecidos são siglas e nomes técnicos (ex.: tcp, dns, json, kubernetes) e palavras comuns "
                      "que faltam no dicionário (ex.: gerenciamento, respectivamente), não erros de conversão.")
    outras.append("Subáreas com pelo menos 500 questões: " +
                  ", ".join(s["nome"] for s in E["por_subarea"].values() if s["questoes"] >= 500) + ".")

    return {
        4: "QuestõesTI-FGV",
        5: CRIADOR,
        6: LOCAL,
        7: f"Versão {VERSAO} - criada em {hoje}",
        8: f"Versão {VERSAO} - sem atualizações desde a criação ({hoje})",
        9: (f"O dataset foi construído a partir de provas de concursos públicos disponibilizadas na web. A fonte foi o "
            f"PCI Concursos (https://www.pciconcursos.com.br/provas/ti/). A coleta foi realizada em setembro de 2026: "
            f"foram selecionadas {g['provas_utilizadas']} provas de cargos de TI da banca FGV, com os respectivos "
            f"gabaritos, baixadas manualmente (o site exige verificação de segurança por prova). Os concursos das "
            f"provas referem-se aos anos de {anos[0]} a {anos[-1]}."),
        10: "Computação / Tecnologia da Informação (questões de concursos públicos)",
        11: "Classificação de Texto (subárea da questão); também serve para resposta a questões de múltipla escolha",
        12: "Português (Brasil), com termos técnicos em inglês",
        13: ("JSON, o dataset está organizado em questões (corpus.json: metadados + lista de questões com enunciado, "
             "alternativas, gabarito, subárea e ano; por_subarea/*.json: questões agrupadas por ano). Também em CSV e "
             "um TXT por questão."),
        14: g["total_tokens"],
        15: g["total_types"],
        16: g["questoes_no_corpus"],
        17: len(E["por_subarea"]),
        18: f"O dataset possui questões de {len(E['por_subarea'])} subáreas da computação: {dist}.",
        19: f"Tamanho médio em tokens de cada questão (enunciado + alternativas): {g['tokens_por_questao']['media']:.2f}".replace(".", ","),
        20: g["maior_questao"]["tokens"],
        21: g["menor_questao"]["tokens"],
        22: qual["proporcao_reconhecida"] if qual else "não calculada (rodar 04b_qualidade_lexica.py)",
        23: tamanho(CORPUS / "corpus.json") + " (corpus.json)",
        24: " ".join(outras),
    }, E


def main():
    V, E = valores()
    wb = load_workbook(RAIZ / "docs" / "templateDatasetCard.xlsx")
    ws = wb.active
    ws["A1"] = "Dataset Card - QuestõesTI-FGV"
    ws["C3"] = "Valor"
    for linha, v in V.items():
        c = ws[f"C{linha}"]
        c.value = v
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws[f"A{linha}"].alignment = Alignment(wrap_text=True, vertical="top")
        ws[f"B{linha}"].alignment = Alignment(wrap_text=True, vertical="top")
        n_linhas = max(len(str(v)) // 48 + 1, len(str(ws[f"B{linha}"].value or "")) // 48 + 1)
        ws.row_dimensions[linha].height = max(ws.row_dimensions[linha].height or 15, 15 * n_linhas)
    ws["C22"].number_format = "0.0%"
    ws["C6"].hyperlink = None                      # no template esta célula é um link de exemplo
    ws["C6"].font = copy(ws["C5"].font)
    ws["C6"].border = copy(ws["C5"].border)
    ws["C6"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["C"].width = 60
    wb.save(RAIZ / "DATASET_CARD.xlsx")

    L = ["# Dataset Card - QuestõesTI-FGV", "", "| Atributo | Valor |", "|---|---|"]
    for linha, v in V.items():
        nome = str(ws[f"A{linha}"].value).strip().rstrip(":").replace("\xa0", "")
        if linha == 22 and isinstance(v, float):
            v = f"{v:.1%}"
        L.append(f"| {nome} | {str(v).replace('|', '/')} |")
    L += ["", "Gráficos e tabelas completas: [corpus/estatisticas/ESTATISTICAS.md](corpus/estatisticas/ESTATISTICAS.md)", ""]
    (RAIZ / "DATASET_CARD.md").write_text("\n".join(L), encoding="utf-8")
    print("DATASET_CARD.xlsx e DATASET_CARD.md gravados")


if __name__ == "__main__":
    main()
